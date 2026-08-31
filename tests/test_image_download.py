"""本文件验证图片网站下载的分批、坏数据隔离、ZIP分类和报告。"""

from pathlib import Path
from threading import Event
import json
import zipfile

import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
import pytest

from src.image_download import (
    ImageDownloadOrchestrator, ImageDownloadRequest, ProductIdSummary,
    extract_and_classify_archive, extract_product_ids,
)
from src.image_site_automation import EdgeImageSiteBackend


class FakeEdgeOptions:
    def __init__(self) -> None:
        self.arguments: list[str] = []
        self.experimental_options: dict[str, object] = {}

    def add_argument(self, value: str) -> None:
        self.arguments.append(value)

    def add_experimental_option(self, name: str, value: object) -> None:
        self.experimental_options[name] = value


class FakeLoginElement:
    def __init__(self) -> None:
        self.values: list[str] = []
        self.clicked = False

    def clear(self) -> None:
        self.values.clear()

    def send_keys(self, value: str) -> None:
        self.values.append(value)

    def click(self) -> None:
        self.clicked = True

    def is_displayed(self) -> bool:
        return True


class FakeTextArea:
    def __init__(self) -> None:
        self.value = "old"

    def clear(self) -> None:
        self.value = ""

    def send_keys(self, value: str) -> None:
        self.value += value


class FakeDisplayedText:
    def __init__(self, text: str) -> None:
        self.text = text

    def is_displayed(self) -> bool:
        return True


class FakeDialog(FakeDisplayedText):
    def __init__(self, text: str) -> None:
        super().__init__(text)
        self.button = FakeLoginElement()

    def find_elements(self, _by, _selector):
        return [self.button]


class FakeDirectDriver:
    def __init__(self, count_text: str = "一共 2 个产品号") -> None:
        self.current_url = "https://fuel-cell.apac.bosch.com/customize/download?token=secret"
        self.count_text = count_text
        self.scripts: list[tuple[str, object]] = []
        self.dialogs: list[FakeDialog] = []

    def get(self, _url: str) -> None:
        pass

    def execute_script(self, script: str, element=None) -> None:
        self.scripts.append((script, element))

    def find_elements(self, _by, selector: str):
        if "一共" in selector:
            return [FakeDisplayedText(self.count_text)]
        if selector == ".el-message-box":
            return self.dialogs
        return []


def _product(index: int) -> str:
    return f"P{index:024d}"


def _tiff_bytes() -> bytes:
    image = np.full((12, 16), 128, dtype=np.uint8)
    ok, encoded = cv2.imencode(".tif", image)
    assert ok
    return encoded.tobytes()


class FakeBackend:
    def __init__(self, bad_product: str = "", bad_code: str = "") -> None:
        self.bad_product = bad_product
        self.bad_code = bad_code
        self.calls: list[tuple[tuple[str, ...], tuple[str, ...]]] = []
        self.closed = False

    def download_batch(
        self, product_ids, image_codes, quality, skip_rework,
        download_dir, stop_event,
    ) -> Path:
        products, codes = tuple(product_ids), tuple(image_codes)
        self.calls.append((products, codes))
        if self.bad_product in products and (not self.bad_code or self.bad_code in codes):
            raise RuntimeError("one image unavailable")
        archive = download_dir / f"result_{len(self.calls)}.zip"
        with zipfile.ZipFile(archive, "w") as bundle:
            for product in products:
                for code in codes:
                    bundle.writestr(f"{product}20260826{code}.tif", _tiff_bytes())
        return archive

    def close(self) -> None:
        self.closed = True


class OmitOnceBackend(FakeBackend):
    def __init__(self, omitted_product: str, omitted_code: str) -> None:
        super().__init__(); self.omitted_product = omitted_product; self.omitted_code = omitted_code

    def download_batch(self, product_ids, image_codes, quality, skip_rework, download_dir, stop_event):
        products, codes = tuple(product_ids), tuple(image_codes)
        self.calls.append((products, codes))
        archive = download_dir / f"result_{len(self.calls)}.zip"
        with zipfile.ZipFile(archive, "w") as bundle:
            for product in products:
                for code in codes:
                    if len(self.calls) == 1 and product == self.omitted_product and code == self.omitted_code:
                        continue
                    bundle.writestr(f"{product}20260826{code}.tif", _tiff_bytes())
        return archive


def _request(tmp_path: Path, codes=("DA", "DE"), batch_size=80) -> ImageDownloadRequest:
    workbook = tmp_path / "source.xlsx"
    Workbook().save(workbook)
    return ImageDownloadRequest(workbook, tmp_path / "output", tuple(codes), "origin", batch_size=batch_size)


def test_extract_product_ids_filters_and_keeps_production_order(tmp_path: Path) -> None:
    valid_a, valid_b = _product(1), _product(2)
    path = tmp_path / "mes.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "MS0310all"
    sheet.append(["Ident No.", "Test Date", "Line", "ST", "SI", "FU", "WP", "Result.Force"])
    sheet.append([valid_b, "2026-08-26 09:00:00", 3003, 10, 1, 1, 6, 1])
    sheet.append([valid_a, "2026-08-26 08:00:00", 3003, 10, 1, 1, 6, 1])
    sheet.append([valid_a, "2026-08-26 10:00:00", 3003, 10, 1, 1, 6, 1])
    sheet.append(["SHORT", "2026-08-26 11:00:00", 3003, 10, 1, 1, 6, 1])
    book.save(path); book.close()
    result = extract_product_ids(path)
    assert result.valid_ids == (valid_a, valid_b)
    assert result.invalid_ids == ("SHORT",)
    assert result.duplicate_count == 1


def test_eight_sheet_workbook_uses_only_four_dedicated_aoi_sheets(tmp_path: Path) -> None:
    path = tmp_path / "eight_sheets.xlsx"
    book = Workbook(); book.remove(book.active)
    headers = ["Ident No.", "Test Date", "Line", "ST", "SI", "FU", "WP"]
    non_aoi = _product(90)
    for name, location in (
        ("MS0310all", (3003, 10, 1, 1, 1)),
        ("MS0320all", (3004, 20, 1, 1, 1)),
        ("MS0335all", (3002, 35, 1, 1, 1)),
    ):
        sheet = book.create_sheet(name); sheet.append(headers)
        sheet.append([non_aoi, "2026-08-26 07:00:00", *location])
    expected: dict[str, tuple[str, ...]] = {}
    for index, (family, sheet_name) in enumerate(
        (("D", "MS03106"), ("E", "MS03206"), ("F", "MS03301"), ("G", "MS03302")), 1,
    ):
        first, second = _product(index * 10 + 1), _product(index * 10 + 2)
        sheet = book.create_sheet(sheet_name.lower() if family == "D" else sheet_name)
        sheet.append(["Ident No.", "Test Date"])
        sheet.append([second, "2026-08-26 09:00:00"])
        sheet.append([first, "2026-08-26 08:00:00"])
        sheet.append([first, "2026-08-26 10:00:00"])
        sheet.append([f"BAD-{family}", "2026-08-26 11:00:00"])
        expected[family] = (first, second)
    package = book.create_sheet("Package")
    package.append(["Unique Part Ident No.", "Packaging Date"])
    package.append([_product(99), "2026-08-26 12:00:00"])
    book.save(path); book.close()

    result = extract_product_ids(path)

    assert result.products_by_family == expected
    assert non_aoi not in result.valid_ids
    assert _product(99) not in result.valid_ids
    assert result.duplicate_counts_by_family == {family: 1 for family in "DEFG"}
    assert result.invalid_ids_by_family == {
        family: (f"BAD-{family}",) for family in "DEFG"
    }
    assert result.sources_by_family["D"] == "ms03106"


def test_missing_dedicated_sheet_falls_back_to_exact_station_location(tmp_path: Path) -> None:
    product = _product(1)
    path = tmp_path / "fallback.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "MS0310all"
    sheet.append(["Ident No.", "Test Date", "Line", "ST", "SI", "FU", "WP"])
    sheet.append([product, "2026-08-26 08:00:00", 3003, 10, 1, 1, 6])
    sheet.append([_product(2), "2026-08-26 08:01:00", 3003, 10, 1, 1, 5])
    book.save(path); book.close()

    result = extract_product_ids(path)

    assert result.products_by_family["D"] == (product,)
    assert result.products_by_family["E"] == ()
    assert result.sources_by_family["D"] == "工站匹配：MS0310all"


def test_82_products_are_split_into_80_and_2(tmp_path: Path) -> None:
    products = tuple(_product(index) for index in range(82))
    backend = FakeBackend()
    result = ImageDownloadOrchestrator(_request(tmp_path), backend).run(ProductIdSummary(products, (), 0))
    assert result.status == "complete"
    assert [len(call[0]) for call in backend.calls] == [80, 2]
    assert result.completed_item_count == 164
    assert backend.closed


def test_multiple_aoi_families_are_downloaded_without_cross_product_requests(tmp_path: Path) -> None:
    d_products = (_product(1), _product(2))
    e_products = (_product(3),)
    summary = ProductIdSummary(
        d_products + e_products, (), 0,
        products_by_family={"D": d_products, "E": e_products, "F": (), "G": ()},
        sources_by_family={"D": "MS03106", "E": "MS03206"},
    )
    backend = FakeBackend()
    result = ImageDownloadOrchestrator(
        _request(tmp_path, codes=("DE", "EE"), batch_size=10), backend,
    ).run(summary)

    assert backend.calls == [(d_products, ("DE",)), (e_products, ("EE",))]
    assert result.completed_item_count == 3
    assert not list(result.output_dir.rglob("products.xlsx"))
    runtime_log = (result.output_dir / "image_download_run.log").read_text(encoding="utf-8")
    assert "直接填写DMC" in runtime_log
    assert "temporary-secret" not in runtime_log
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["version"] == 2
    assert manifest["product_ids_by_family"] == {"D": list(d_products), "E": list(e_products)}
    report = load_workbook(result.summary_path, read_only=True)
    assert "AOI Stations" in report.sheetnames
    report.close()


def test_selected_family_without_aoi_products_is_rejected_before_backend_call(tmp_path: Path) -> None:
    d_product = _product(1)
    summary = ProductIdSummary(
        (d_product,), (), 0,
        products_by_family={"D": (d_product,), "E": (), "F": (), "G": ()},
    )
    backend = FakeBackend()

    with pytest.raises(ValueError, match="MS03206"):
        ImageDownloadOrchestrator(
            _request(tmp_path, codes=("DE", "EE"), batch_size=10), backend,
        ).run(summary)

    assert backend.calls == []


def test_bad_product_and_code_are_isolated_while_good_images_survive(tmp_path: Path) -> None:
    products = tuple(_product(index) for index in range(8))
    bad = products[3]
    backend = FakeBackend(bad, "DE")
    result = ImageDownloadOrchestrator(
        _request(tmp_path, batch_size=10), backend,
    ).run(ProductIdSummary(products, (), 0))
    assert result.status == "partial"
    assert (result.output_dir / "images" / "5S" / f"{bad}20260826DA.tif").is_file()
    assert not (result.output_dir / "images" / "5S" / f"{bad}20260826DE.tif").exists()
    assert any(issue.product_id == bad and issue.image_code == "DE" for issue in result.issues)
    for product in products:
        assert (result.output_dir / "images" / "5S" / f"{product}20260826DA.tif").is_file()


def test_corrupt_tiff_is_quarantined(tmp_path: Path) -> None:
    product = _product(1)
    archive = tmp_path / "bad.zip"
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.writestr(f"{product}20260826DA.tif", b"not-a-tiff")
    outcome = extract_and_classify_archive(archive, tmp_path, [product], ["DA"])
    assert not outcome.found_items
    assert outcome.issues[0].category == "corrupt_tiff"
    assert any((tmp_path / "_quarantine").iterdir())


def test_resize_suffix_is_normalized_for_existing_image_index(tmp_path: Path) -> None:
    product = _product(1)
    archive = tmp_path / "resize.zip"
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.writestr(f"{product}20260826DA_resize.tif", _tiff_bytes())
    outcome = extract_and_classify_archive(archive, tmp_path, [product], ["DA"])
    assert outcome.found_items == {(product, "DA")}
    assert (tmp_path / "images" / "5S" / f"{product}20260826DA.tif").is_file()


def test_missing_image_in_successful_zip_is_retried_individually(tmp_path: Path) -> None:
    products = tuple(_product(index) for index in range(4))
    backend = OmitOnceBackend(products[2], "DE")
    result = ImageDownloadOrchestrator(
        _request(tmp_path, batch_size=10), backend,
    ).run(ProductIdSummary(products, (), 0))
    assert result.status == "complete"
    assert backend.calls[-1] == ((products[2],), ("DE",))
    assert result.completed_item_count == 8


def test_retry_manifest_only_downloads_failed_item(tmp_path: Path) -> None:
    products = tuple(_product(index) for index in range(4))
    bad = products[1]
    first = ImageDownloadOrchestrator(
        _request(tmp_path, batch_size=10), FakeBackend(bad, "DE"),
    ).run(ProductIdSummary(products, (), 0))
    retry_request = ImageDownloadRequest(
        first.manifest_path.parent.parent / "source.xlsx",
        first.output_dir, ("DA", "DE"), "origin", batch_size=10,
        retry_manifest=first.manifest_path,
    )
    backend = FakeBackend()
    retried = ImageDownloadOrchestrator(retry_request, backend).run(ProductIdSummary(products, (), 0))
    assert retried.status == "complete"
    assert backend.calls == [((bad,), ("DE",))]
    assert retried.product_count == 4
    assert retried.completed_item_count == 8


def test_v1_retry_manifest_retries_exact_sparse_pairs(tmp_path: Path) -> None:
    products = (_product(1), _product(2))
    request = _request(tmp_path, codes=("DA", "DE"), batch_size=10)
    manifest_path = tmp_path / "v1_manifest.json"
    manifest_path.write_text(json.dumps({
        "version": 1,
        "workbook": str(request.workbook_path),
        "image_codes": ["DA", "DE"],
        "product_ids": list(products),
        "completed_items": [[products[0], "DE"], [products[1], "DA"]],
        "failed_items": [
            {"product_id": products[0], "image_code": "DA", "category": "missing_image"},
            {"product_id": products[1], "image_code": "DE", "category": "missing_image"},
        ],
    }), encoding="utf-8")
    retry_request = ImageDownloadRequest(
        request.workbook_path, request.output_root, request.image_codes, request.quality,
        batch_size=10, retry_manifest=manifest_path,
    )
    backend = FakeBackend()

    result = ImageDownloadOrchestrator(retry_request, backend).run(
        ProductIdSummary(products, (), 0),
    )

    assert backend.calls == [((products[0],), ("DA",)), ((products[1],), ("DE",))]
    assert result.status == "complete"
    assert result.completed_item_count == 4


def test_product_limit_and_credentials_are_not_persisted(tmp_path: Path) -> None:
    request = _request(tmp_path)
    secret_request = ImageDownloadRequest(
        request.workbook_path, request.output_root, request.image_codes, request.quality,
        username="ldap-user", password="super-secret-password", batch_size=80,
    )
    result = ImageDownloadOrchestrator(secret_request, FakeBackend()).run(
        ProductIdSummary((_product(1),), (), 0)
    )
    manifest = result.manifest_path.read_text(encoding="utf-8")
    assert "ldap-user" not in manifest
    assert "super-secret-password" not in manifest

    too_large = ImageDownloadRequest(
        request.workbook_path, request.output_root, request.image_codes, request.quality,
        batch_size=101,
    )
    with pytest.raises(ValueError, match="10到100"):
        too_large.validate()
    backend = EdgeImageSiteBackend(tmp_path)
    with pytest.raises(ValueError, match="最多接受100"):
        backend.download_batch(
            [_product(index) for index in range(101)], ["DA"], "origin", True,
            tmp_path, Event(),
        )


def test_edge_options_use_dedicated_persistent_profile(tmp_path: Path) -> None:
    profile_dir = tmp_path / "edge-image-profile"
    download_dir = tmp_path / "downloads"
    download_dir.mkdir()
    backend = EdgeImageSiteBackend(tmp_path, profile_dir=profile_dir)

    options = backend._build_options(FakeEdgeOptions, download_dir)

    assert "--inprivate" not in options.arguments
    assert f"--user-data-dir={profile_dir.resolve()}" in options.arguments
    assert "--profile-directory=Default" in options.arguments
    assert profile_dir.is_dir()
    assert options.experimental_options["prefs"]["download.default_directory"] == str(
        download_dir.resolve()
    )


def test_direct_entry_fills_newline_dmcs_and_verifies_site_count(tmp_path: Path) -> None:
    messages: list[str] = []
    clicked: list[str] = []
    checked: list[tuple[str, bool]] = []
    textarea = FakeTextArea()
    backend = EdgeImageSiteBackend(tmp_path, log=messages.append, profile_dir=tmp_path / "profile")
    backend.driver = FakeDirectDriver()
    backend._By = type("FakeBy", (), {"CSS_SELECTOR": "css", "XPATH": "xpath"})
    backend._page_ready = lambda: True
    backend._click_xpath = lambda xpath, root=None: clicked.append(xpath)
    backend._wait_visible = lambda by, selector, timeout=30: textarea
    backend._set_checkbox = lambda label, value, root=None: checked.append((label, value))
    products = (_product(1), _product(2))

    backend._configure_direct_page(products, ("DA", "DE"), "origin", True)

    assert textarea.value == "\n".join(products)
    assert any("not(contains(.,'导入'))" in xpath for xpath in clicked)
    assert any("只选择原图" in xpath for xpath in clicked)
    assert ("DA", True) in checked and ("DE", True) in checked
    assert ("EA", False) in checked and ("去除7层返工站", True) in checked
    assert any("new Event('input'" in script for script, _element in backend.driver.scripts)
    assert any("提交2个，网站识别2个" in message for message in messages)
    assert all("token=secret" not in message for message in messages)


def test_direct_entry_rejects_recognized_count_mismatch(tmp_path: Path, monkeypatch) -> None:
    backend = EdgeImageSiteBackend(tmp_path, profile_dir=tmp_path / "profile")
    backend.driver = FakeDirectDriver("一共识别 1 个产品号")
    backend._By = type("FakeBy", (), {"XPATH": "xpath"})
    ticks = iter((0.0, 0.0, 2.0))
    monkeypatch.setattr("src.image_site_automation.time.monotonic", lambda: next(ticks))
    monkeypatch.setattr("src.image_site_automation.time.sleep", lambda _seconds: None)

    with pytest.raises(RuntimeError, match="DMC解析数量不一致：提交2个，网站识别1个"):
        backend._wait_recognized_count(2, timeout=1)


@pytest.mark.parametrize(
    ("missing_control", "expected_message"),
    (("tab", "未找到直接输入页签"), ("textarea", "未找到可见的产品号文本框")),
)
def test_direct_entry_reports_unavailable_controls(
    tmp_path: Path, missing_control: str, expected_message: str,
) -> None:
    backend = EdgeImageSiteBackend(tmp_path, profile_dir=tmp_path / "profile")
    backend.driver = FakeDirectDriver()
    backend._By = type("FakeBy", (), {"CSS_SELECTOR": "css", "XPATH": "xpath"})
    backend._page_ready = lambda: True
    if missing_control == "tab":
        backend._click_xpath = lambda _xpath: (_ for _ in ()).throw(TimeoutError("missing"))
    else:
        backend._click_xpath = lambda _xpath: None
        backend._wait_visible = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            TimeoutError("missing")
        )

    with pytest.raises(RuntimeError, match=expected_message):
        backend._configure_direct_page((_product(1),), ("DA",), "origin", True)


def test_only_semantic_error_dialogs_abort_download(tmp_path: Path) -> None:
    backend = EdgeImageSiteBackend(tmp_path, profile_dir=tmp_path / "profile")
    driver = FakeDirectDriver()
    backend.driver = driver
    backend._By = type("FakeBy", (), {"CSS_SELECTOR": "css", "XPATH": "xpath"})
    backend._find_first = lambda _selectors: None
    neutral = FakeDialog("文件正在生成，请稍候")
    driver.dialogs = [neutral]

    assert backend._website_error_text() == ""
    backend._confirm_neutral_dialog()
    assert neutral.button.clicked

    driver.dialogs = [FakeDialog("下载过程中发生错误，请重试")]
    assert backend._website_error_text() == "下载过程中发生错误，请重试"

    backend._find_first = lambda _selectors: FakeDisplayedText("服务器处理失败")
    assert backend._website_error_text() == "服务器处理失败"


def test_default_edge_profile_is_stable_and_application_owned(tmp_path: Path, monkeypatch) -> None:
    app_data = tmp_path / "MEA5SDefectAnalysis"
    monkeypatch.setattr("src.image_site_automation.user_data_dir", lambda: app_data)

    first = EdgeImageSiteBackend(tmp_path)
    second = EdgeImageSiteBackend(tmp_path)

    assert first.profile_dir == app_data / "edge-image-profile"
    assert second.profile_dir == first.profile_dir


def test_wia_login_is_reported_as_manual_and_times_out_clearly(tmp_path: Path, monkeypatch) -> None:
    messages: list[str] = []
    backend = EdgeImageSiteBackend(
        tmp_path, password="temporary-secret", log=messages.append,
        login_wait_seconds=1, profile_dir=tmp_path / "profile",
    )
    backend._page_ready = lambda: False
    backend._current_url = lambda: "https://stfs.bosch.com/adfs/ls/wia?wa=wsignin1.0"
    ticks = iter((0.0, 0.0, 2.0))
    monkeypatch.setattr("src.image_site_automation.time.monotonic", lambda: next(ticks))
    monkeypatch.setattr("src.image_site_automation.time.sleep", lambda _seconds: None)

    with pytest.raises(TimeoutError, match="Windows集成认证超时"):
        backend._complete_login()

    assert any("系统窗口" in message and "手动完成" in message for message in messages)
    assert backend.password == ""


def test_web_ldap_form_is_filled_and_login_completion_is_detected(tmp_path: Path, monkeypatch) -> None:
    messages: list[str] = []
    username = FakeLoginElement()
    password = FakeLoginElement()
    button = FakeLoginElement()
    backend = EdgeImageSiteBackend(
        tmp_path, username="ldap-user", password="temporary-secret",
        log=messages.append, profile_dir=tmp_path / "profile",
    )
    ready = iter((False, True))
    elements = iter((username, password, button))
    backend._page_ready = lambda: next(ready)
    backend._current_url = lambda: "https://fuel-cell.apac.bosch.com/login"
    backend._find_first = lambda _selectors: next(elements)
    backend._By = type(
        "FakeBy", (), {"ID": "id", "NAME": "name", "CSS_SELECTOR": "css", "XPATH": "xpath"},
    )
    monkeypatch.setattr("src.image_site_automation.time.sleep", lambda _seconds: None)

    backend._complete_login()

    assert username.values == ["ldap-user"]
    assert password.values == ["temporary-secret"]
    assert button.clicked
    assert any("网页LDAP登录表单" in message for message in messages)
    assert messages[-1] == "图片网站登录成功"
