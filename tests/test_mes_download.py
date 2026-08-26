"""本文件验证MES下载服务的参数检查、任务编排和工作簿交付。"""

from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from src import mes_download
from src.mes_download import MesDownloadRequest, download_mes_workbook


def test_mes_request_rejects_invalid_range_and_model(tmp_path: Path) -> None:
    now = datetime(2026, 8, 26, 8)
    with pytest.raises(ValueError, match="结束时间"):
        MesDownloadRequest(now, now, tmp_path).validate()
    with pytest.raises(ValueError, match="10位"):
        MesDownloadRequest(now, now + timedelta(hours=1), tmp_path, "123").validate()


def test_mes_service_runs_queries_and_returns_unique_range_workbook(
    monkeypatch, tmp_path: Path,
) -> None:
    calls: list[str] = []

    class FakeExcelProcessor:
        @staticmethod
        def create_failure_rate_excel(folder: str, date_name: str) -> str:
            path = Path(folder) / f"Failure rate_{date_name}.xlsx"
            workbook = Workbook()
            workbook.save(path)
            workbook.close()
            return str(path)

        @staticmethod
        def process_all_sheets(path: str) -> None:
            workbook = load_workbook(path)
            workbook.active["A1"] = "processed"
            workbook.save(path)
            workbook.close()
            calls.append("processed")

    class FakeAutomation:
        def __init__(self, download_dir, **kwargs) -> None:
            self.stop_event = None
            self.kwargs = kwargs
            calls.append("created")

        def start_browser(self) -> None:
            calls.append("browser")

        def open_portal(self) -> None:
            calls.append("portal")

        def navigate_to_quality_data(self) -> None:
            calls.append("quality")

        def execute_query(self, task, *_args) -> bool:
            calls.append(task["name"])
            return True

        def close(self, kill_browser=False) -> None:
            calls.append(f"closed:{kill_browser}")

    monkeypatch.setattr(
        mes_download, "_load_legacy_modules",
        lambda _root: (
            SimpleNamespace(QUERY_TASKS=[{"name": "MS0310all"}, {"name": "VI"}]),
            FakeExcelProcessor,
            SimpleNamespace(OISAutomation=FakeAutomation),
        ),
    )
    request = MesDownloadRequest(
        datetime(2026, 8, 26, 8), datetime(2026, 8, 26, 12), tmp_path,
        username="user", password="secret",
    )
    path = download_mes_workbook(tmp_path, request)

    assert path.name == "Failure rate_20260826_0800-20260826_1200.xlsx"
    assert path.is_file()
    assert load_workbook(path, read_only=True).active["A1"].value == "processed"
    assert calls == [
        "created", "browser", "portal", "quality", "MS0310all", "VI",
        "processed", "closed:False",
    ]

