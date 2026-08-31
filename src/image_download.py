"""本文件编排公司图片网站批量下载、坏数据隔离、解压分类和结果报告。"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from threading import Event
from typing import Callable, Protocol, Sequence
import csv
import hashlib
import json
import re
import shutil
import zipfile

import cv2
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

from src.station_sources import parse_image_filename


ALL_IMAGE_CODES = (
    "DA", "DB", "DC", "DE", "DX", "DY",
    "EA", "EB", "EC", "EE", "EX", "EY",
    "FA", "FB", "FC", "FE", "FX", "FY",
    "GA", "GB", "GC", "GE", "GX", "GY",
)
PRIMARY_IMAGE_CODES = ("DA", "DE", "EA", "EE", "FA", "FE", "GA", "GE")
CODE_SCOPE = {"D": "5S", "E": "5X", "F": "7S", "G": "7X"}
MIN_FREE_BYTES = 2 * 1024 * 1024 * 1024


ProgressCallback = Callable[[int, str], None]
LogCallback = Callable[[str], None]


@dataclass(frozen=True)
class AoiDownloadGroup:
    family: str
    scope: str
    sheet_name: str
    station_id: str
    station_name: str


AOI_DOWNLOAD_GROUPS = (
    AoiDownloadGroup("D", "5S", "MS03106", "35_5s_aoi", "5S AOI"),
    AoiDownloadGroup("E", "5X", "MS03206", "57_5x_aoi", "5X AOI"),
    AoiDownloadGroup("F", "7S", "MS03301", "conveyor_7s_aoi", "7S AOI"),
    AoiDownloadGroup("G", "7X", "MS03302", "conveyor_7x_aoi", "7X AOI"),
)
AOI_GROUP_BY_FAMILY = {group.family: group for group in AOI_DOWNLOAD_GROUPS}


@dataclass(frozen=True)
class ProductIdSummary:
    valid_ids: tuple[str, ...]
    invalid_ids: tuple[str, ...]
    duplicate_count: int
    products_by_family: dict[str, tuple[str, ...]] = field(default_factory=dict)
    invalid_ids_by_family: dict[str, tuple[str, ...]] = field(default_factory=dict)
    duplicate_counts_by_family: dict[str, int] = field(default_factory=dict)
    sources_by_family: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class ImageDownloadRequest:
    workbook_path: Path
    output_root: Path
    image_codes: tuple[str, ...]
    quality: str
    skip_rework: bool = True
    username: str = ""
    password: str = ""
    batch_size: int = 80
    retry_manifest: Path | None = None

    def validate(self) -> None:
        if not self.workbook_path.is_file():
            raise FileNotFoundError(f"MES工作簿不存在：{self.workbook_path}")
        if not self.image_codes:
            raise ValueError("请至少选择一个图片代码")
        unknown = sorted(set(self.image_codes) - set(ALL_IMAGE_CODES))
        if unknown:
            raise ValueError("未知图片代码：" + "、".join(unknown))
        if self.quality not in {"origin", "resize"}:
            raise ValueError("必须选择原图或压缩图")
        if not 10 <= self.batch_size <= 100:
            raise ValueError("每批产品号必须在10到100之间")


@dataclass
class ImageDownloadIssue:
    product_id: str
    image_code: str
    stage: str
    category: str
    message: str
    retry_count: int = 0
    status: str = "failed"


@dataclass
class ImageDownloadResult:
    status: str
    output_dir: Path
    image_roots: dict[str, Path]
    product_count: int
    completed_item_count: int
    issues: list[ImageDownloadIssue]
    manifest_path: Path
    summary_path: Path


class ImageSiteBackend(Protocol):
    def download_batch(
        self,
        product_ids: Sequence[str],
        image_codes: Sequence[str],
        quality: str,
        skip_rework: bool,
        download_dir: Path,
        stop_event: Event,
    ) -> Path:
        """操作网站并返回下载完成的ZIP路径。"""

    def close(self) -> None:
        """关闭浏览器和驱动。"""


def _normalized_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _timestamp(value: object) -> pd.Timestamp:
    text = str(value or "").strip()
    if not text:
        return pd.NaT
    year_first = bool(re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", text))
    return pd.to_datetime(text, errors="coerce", yearfirst=year_first, dayfirst=not year_first)


def _read_aoi_sheet_records(workbook, expected_name: str) -> tuple[list[tuple[str, pd.Timestamp, int]], str]:
    """按名称读取专用AOI页签中的Ident No.记录，名称匹配不区分大小写。"""
    actual_name = next(
        (name for name in workbook.sheetnames if name.strip().casefold() == expected_name.casefold()),
        "",
    )
    if not actual_name:
        return [], ""
    rows = list(workbook[actual_name].iter_rows(values_only=True))
    starts = [
        index for index, row in enumerate(rows)
        if row and any(_normalized_header(value) == "identno" for value in row)
    ]
    records: list[tuple[str, pd.Timestamp, int]] = []
    for position, start in enumerate(starts):
        end = starts[position + 1] if position + 1 < len(starts) else len(rows)
        normalized = [_normalized_header(value) for value in rows[start]]
        ident_index = normalized.index("identno")
        date_index = normalized.index("testdate") if "testdate" in normalized else -1
        for row_index in range(start + 1, end):
            row = rows[row_index]
            if ident_index >= len(row):
                continue
            product_id = str(row[ident_index] or "").strip().upper()
            if not product_id:
                continue
            date_value = row[date_index] if 0 <= date_index < len(row) else None
            records.append((product_id, _timestamp(date_value), row_index + 1))
    return records, actual_name


def _summarize_records(
    records: Sequence[tuple[str, pd.Timestamp, int]],
) -> tuple[tuple[str, ...], tuple[str, ...], int]:
    ordered = sorted(
        records,
        key=lambda item: (
            pd.isna(item[1]), item[1] if pd.notna(item[1]) else pd.Timestamp.max, item[2],
        ),
    )
    valid: list[str] = []
    invalid: list[str] = []
    seen: set[str] = set()
    duplicate_count = 0
    for product_id, _test_date, _row_number in ordered:
        if len(product_id) != 25:
            if product_id not in invalid:
                invalid.append(product_id)
            continue
        if product_id in seen:
            duplicate_count += 1
            continue
        seen.add(product_id)
        valid.append(product_id)
    return tuple(valid), tuple(invalid), duplicate_count


def extract_product_ids(workbook_path: Path) -> ProductIdSummary:
    """按四个AOI工站提取产品号，优先使用专用页签并按工站位置兜底。"""
    from src.station_sources import load_station_catalog
    from src.station_workbook import load_station_workbook

    resolved = workbook_path.resolve()
    if not resolved.is_file():
        raise FileNotFoundError(f"Excel工作簿不存在：{resolved}")
    exact_records: dict[str, list[tuple[str, pd.Timestamp, int]]] = {}
    exact_sources: dict[str, str] = {}
    book = load_workbook(
        resolved, read_only=True, data_only=True,
        keep_vba=resolved.suffix.lower() == ".xlsm",
    )
    try:
        for group in AOI_DOWNLOAD_GROUPS:
            records, source = _read_aoi_sheet_records(book, group.sheet_name)
            exact_records[group.family] = records
            exact_sources[group.family] = source
    finally:
        book.close()

    missing_families = {
        group.family for group in AOI_DOWNLOAD_GROUPS if not exact_records[group.family]
    }
    fallback_events = None
    if missing_families:
        project_root = Path(__file__).resolve().parents[1]
        catalog = load_station_catalog(project_root / "config" / "stations.yaml")
        try:
            fallback_events = load_station_workbook(resolved, catalog).events.copy()
        except ValueError as exc:
            if "没有可识别的工站记录" not in str(exc):
                raise
            fallback_events = pd.DataFrame(columns=[
                "station_id", "dmc_raw", "test_date", "source_row", "source_sheet",
            ])

    products_by_family: dict[str, tuple[str, ...]] = {}
    invalid_by_family: dict[str, tuple[str, ...]] = {}
    duplicates_by_family: dict[str, int] = {}
    sources_by_family: dict[str, str] = {}
    for group in AOI_DOWNLOAD_GROUPS:
        records = exact_records[group.family]
        source = exact_sources[group.family]
        if not records and fallback_events is not None:
            station_events = fallback_events[
                fallback_events["station_id"].astype(str).eq(group.station_id)
            ]
            records = [
                (str(row.dmc_raw or "").strip().upper(), row.test_date, int(row.source_row))
                for row in station_events.itertuples()
                if str(row.dmc_raw or "").strip()
            ]
            fallback_sources = tuple(dict.fromkeys(
                str(value) for value in station_events.get("source_sheet", pd.Series(dtype=str))
                if str(value).strip()
            ))
            source = "工站匹配：" + "、".join(fallback_sources) if fallback_sources else "未找到"
        valid, invalid, duplicates = _summarize_records(records)
        products_by_family[group.family] = valid
        invalid_by_family[group.family] = invalid
        duplicates_by_family[group.family] = duplicates
        sources_by_family[group.family] = source or "未找到"

    valid_ids = tuple(dict.fromkeys(
        product_id
        for group in AOI_DOWNLOAD_GROUPS
        for product_id in products_by_family[group.family]
    ))
    invalid_ids = tuple(dict.fromkeys(
        product_id
        for group in AOI_DOWNLOAD_GROUPS
        for product_id in invalid_by_family[group.family]
    ))
    return ProductIdSummary(
        valid_ids,
        invalid_ids,
        sum(duplicates_by_family.values()),
        products_by_family,
        invalid_by_family,
        duplicates_by_family,
        sources_by_family,
    )


def _safe_member_name(name: str) -> str:
    normalized = name.replace("\\", "/")
    if normalized.startswith("/") or ".." in Path(normalized).parts:
        raise ValueError(f"ZIP包含不安全路径：{name}")
    return Path(normalized).name


def _valid_tiff(data: bytes) -> bool:
    if len(data) < 64 or data[:4] not in {b"II*\x00", b"MM\x00*"}:
        return False
    decoded = cv2.imdecode(np.frombuffer(data, dtype=np.uint8), cv2.IMREAD_UNCHANGED)
    return decoded is not None and decoded.size > 0


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


@dataclass
class _ArchiveOutcome:
    found_items: set[tuple[str, str]] = field(default_factory=set)
    issues: list[ImageDownloadIssue] = field(default_factory=list)
    extracted_count: int = 0


def extract_and_classify_archive(
    archive: Path,
    output_dir: Path,
    expected_products: Sequence[str],
    expected_codes: Sequence[str],
) -> _ArchiveOutcome:
    """校验ZIP和TIFF，并把图片安全分类到四个分析范围。"""
    expected_product_set = set(expected_products)
    expected_code_set = set(expected_codes)
    outcome = _ArchiveOutcome()
    quarantine = output_dir / "_quarantine"
    quarantine.mkdir(parents=True, exist_ok=True)
    try:
        with zipfile.ZipFile(archive) as bundle:
            broken = bundle.testzip()
            if broken:
                raise zipfile.BadZipFile(f"ZIP成员损坏：{broken}")
            for info in bundle.infolist():
                if info.is_dir():
                    continue
                filename = _safe_member_name(info.filename)
                data = bundle.read(info)
                suffix = Path(filename).suffix.lower()
                logical_filename = filename
                if Path(filename).stem.lower().endswith("_resize"):
                    logical_filename = Path(filename).stem[:-7] + Path(filename).suffix
                parsed = parse_image_filename(Path(logical_filename))
                if suffix not in {".tif", ".tiff"} or parsed is None:
                    target = quarantine / filename
                    target.write_bytes(data)
                    outcome.issues.append(ImageDownloadIssue(
                        "", "", "extract", "unrecognized_filename",
                        f"文件名无法识别：{filename}", status="quarantined",
                    ))
                    continue
                product_id, code = parsed.dmc_raw.upper(), parsed.image_code.upper()
                if product_id not in expected_product_set or code not in expected_code_set:
                    target = quarantine / filename
                    target.write_bytes(data)
                    outcome.issues.append(ImageDownloadIssue(
                        product_id, code, "extract", "unexpected_image",
                        f"ZIP包含未请求的图片：{filename}", status="quarantined",
                    ))
                    continue
                if not _valid_tiff(data):
                    target = quarantine / filename
                    target.write_bytes(data)
                    outcome.issues.append(ImageDownloadIssue(
                        product_id, code, "validate", "corrupt_tiff",
                        f"TIFF无法解码：{filename}", status="retryable",
                    ))
                    continue
                scope = CODE_SCOPE[code[0]]
                target_dir = output_dir / "images" / scope
                target_dir.mkdir(parents=True, exist_ok=True)
                # 压缩图的网站文件名带 _resize；分类时去掉该标记，以复用现有现场命名解析。
                target = target_dir / logical_filename
                if target.exists():
                    if _digest(target.read_bytes()) == _digest(data):
                        outcome.found_items.add((product_id, code))
                        continue
                    conflict = quarantine / f"duplicate_{filename}"
                    conflict.write_bytes(data)
                    outcome.issues.append(ImageDownloadIssue(
                        product_id, code, "extract", "duplicate_image",
                        f"同名图片内容不同：{filename}", status="quarantined",
                    ))
                    continue
                target.write_bytes(data)
                outcome.found_items.add((product_id, code))
                outcome.extracted_count += 1
    except zipfile.BadZipFile as exc:
        raise ValueError(f"下载文件不是有效ZIP：{archive.name}；{exc}") from exc
    return outcome


class ImageDownloadOrchestrator:
    def __init__(
        self,
        request: ImageDownloadRequest,
        backend: ImageSiteBackend,
        stop_event: Event | None = None,
        log: LogCallback | None = None,
        progress: ProgressCallback | None = None,
    ) -> None:
        self.request = request
        self.backend = backend
        self.stop_event = stop_event or Event()
        self._external_log = log or (lambda _message: None)
        self.progress = progress or (lambda _value, _message: None)
        self.issues: list[ImageDownloadIssue] = []
        self.completed_items: set[tuple[str, str]] = set()
        self._operation_index = 0
        self.output_dir = request.output_root
        self.temp_dir = self.output_dir / "_temporary_batches"
        self.manifest_path = self.output_dir / "image_download_manifest.json"
        self._all_product_ids: list[str] = []
        self._products_by_family: dict[str, tuple[str, ...]] = {}
        self._sources_by_family: dict[str, str] = {}
        self._planned_items: set[tuple[str, str]] = set()
        self.runtime_log_path = self.output_dir / "image_download_run.log"
        self.log = self._log

    def _log(self, message: str) -> None:
        self._external_log(message)
        try:
            self.output_dir.mkdir(parents=True, exist_ok=True)
            timestamp = datetime.now().astimezone().isoformat(timespec="seconds")
            with self.runtime_log_path.open("a", encoding="utf-8") as handle:
                handle.write(f"{timestamp} | {message}\n")
        except OSError:
            pass

    def _check_cancelled(self) -> None:
        if self.stop_event.is_set():
            raise InterruptedError("图片下载已取消")

    def _check_disk(self) -> None:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        if shutil.disk_usage(self.output_dir).free < MIN_FREE_BYTES:
            raise OSError("图片保存目录剩余空间不足2 GB")

    def _record_manifest(self) -> None:
        payload = {
            "version": 2,
            "updated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "workbook": str(self.request.workbook_path.resolve()),
            "quality": self.request.quality,
            "skip_rework": self.request.skip_rework,
            "image_codes": list(self.request.image_codes),
            "batch_size": self.request.batch_size,
            "product_ids": list(self._all_product_ids),
            "product_ids_by_family": {
                family: list(product_ids)
                for family, product_ids in self._products_by_family.items()
            },
            "sources_by_family": self._sources_by_family,
            "planned_items": [list(item) for item in sorted(self._planned_items)],
            "completed_items": [list(item) for item in sorted(self.completed_items)],
            "failed_items": [
                {"product_id": issue.product_id, "image_code": issue.image_code,
                 "category": issue.category, "message": issue.message}
                for issue in self.issues if issue.status == "failed" and issue.product_id
            ],
        }
        self.manifest_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8",
        )

    def _checkpoint(self) -> None:
        if self._all_product_ids:
            self._record_manifest()

    def _attempt_download(
        self, product_ids: Sequence[str], image_codes: Sequence[str], retry_count: int,
    ) -> Path:
        self._check_cancelled()
        self._check_disk()
        self._operation_index += 1
        label = f"batch_{self._operation_index:04d}"
        batch_dir = self.temp_dir / label
        batch_dir.mkdir(parents=True, exist_ok=True)
        self.log(
            f"{label}：直接填写DMC，{len(product_ids)}个产品，代码 {'+'.join(image_codes)}"
            + (f"，第{retry_count}次重试" if retry_count else "")
        )
        return self.backend.download_batch(
            product_ids, image_codes, self.request.quality, self.request.skip_rework,
            batch_dir, self.stop_event,
        )

    def _download_or_isolate(
        self, product_ids: Sequence[str], image_codes: Sequence[str], verify_retry: int = 0,
    ) -> None:
        self._check_cancelled()
        requested = {
            (product_id, code) for product_id in product_ids for code in image_codes
        } - self.completed_items
        if not requested:
            return
        active_products = [item for item in product_ids if any((item, c) in requested for c in image_codes)]
        active_codes = [code for code in image_codes if any((p, code) in requested for p in active_products)]
        last_error: Exception | None = None
        archive: Path | None = None
        for retry in range(3):
            try:
                archive = self._attempt_download(active_products, active_codes, retry)
                if not archive.is_file() or archive.stat().st_size == 0:
                    raise FileNotFoundError("网站未生成ZIP下载文件")
                break
            except InterruptedError:
                raise
            except Exception as exc:
                last_error = exc
                self.log(f"下载失败：{exc}")
        if archive is None:
            if len(active_products) > 1:
                middle = len(active_products) // 2
                self.log(f"自动拆分产品：{len(active_products)} → {middle}+{len(active_products)-middle}")
                self._download_or_isolate(active_products[:middle], active_codes, verify_retry)
                self._download_or_isolate(active_products[middle:], active_codes, verify_retry)
                return
            if len(active_codes) > 1:
                middle = len(active_codes) // 2
                self.log(f"隔离产品 {active_products[0]} 的图片代码")
                self._download_or_isolate(active_products, active_codes[:middle], verify_retry)
                self._download_or_isolate(active_products, active_codes[middle:], verify_retry)
                return
            self.issues.append(ImageDownloadIssue(
                active_products[0], active_codes[0], "download", "website_unavailable",
                str(last_error or "网站未返回文件"), retry_count=2, status="failed",
            ))
            self._checkpoint()
            return

        try:
            outcome = extract_and_classify_archive(
                archive, self.output_dir, active_products, active_codes,
            )
        except Exception as exc:
            self.log(f"ZIP校验失败，开始隔离：{exc}")
            if len(active_products) > 1:
                middle = len(active_products) // 2
                self._download_or_isolate(active_products[:middle], active_codes, verify_retry)
                self._download_or_isolate(active_products[middle:], active_codes, verify_retry)
            elif len(active_codes) > 1:
                middle = len(active_codes) // 2
                self._download_or_isolate(active_products, active_codes[:middle], verify_retry)
                self._download_or_isolate(active_products, active_codes[middle:], verify_retry)
            else:
                self.issues.append(ImageDownloadIssue(
                    active_products[0], active_codes[0], "extract", "invalid_zip",
                    str(exc), retry_count=2, status="failed",
                ))
                self._checkpoint()
            return

        self.completed_items.update(outcome.found_items)
        self.issues.extend(outcome.issues)
        self._checkpoint()
        missing = requested - outcome.found_items
        try:
            archive.unlink()
        except OSError:
            pass
        # A successful ZIP may still omit or corrupt one image. Retry those items alone.
        retryable = {
            (issue.product_id, issue.image_code)
            for issue in outcome.issues if issue.status == "retryable"
        }
        for product_id, code in sorted(missing | retryable):
            if (product_id, code) in self.completed_items:
                continue
            if verify_retry < 1:
                self._download_or_isolate([product_id], [code], verify_retry + 1)
            if (product_id, code) in self.completed_items:
                for issue in self.issues:
                    if issue.product_id == product_id and issue.image_code == code and issue.status == "retryable":
                        issue.status = "recovered"
                continue
            already_failed = any(
                issue.product_id == product_id and issue.image_code == code and issue.status == "failed"
                for issue in self.issues
            )
            if not already_failed:
                self.issues.append(ImageDownloadIssue(
                    product_id, code, "verify", "missing_image",
                    "ZIP中未找到可用的所选图片", retry_count=verify_retry + 1, status="failed",
                ))
        self._checkpoint()

    def _write_reports(self) -> Path:
        issue_path = self.output_dir / "image_download_issues.csv"
        columns = ["product_id", "image_code", "stage", "category", "message", "retry_count", "status"]
        with issue_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns)
            writer.writeheader()
            for issue in self.issues:
                writer.writerow(asdict(issue))
        summary_path = self.output_dir / "image_download_summary.xlsx"
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["项目", "数量"])
        summary.append(["唯一产品号", len(self._all_product_ids)])
        summary.append(["计划产品-代码项", len(self._planned_items)])
        summary.append(["已下载产品-代码项", len(self.completed_items)])
        summary.append(["异常项", len([i for i in self.issues if i.status == "failed"])])
        stations = workbook.create_sheet("AOI Stations")
        stations.append(["图片范围", "AOI页签", "数据来源", "产品数", "计划项", "完成项"])
        for group in AOI_DOWNLOAD_GROUPS:
            family = group.family
            stations.append([
                group.scope,
                group.sheet_name,
                self._sources_by_family.get(family, ""),
                len(self._products_by_family.get(family, ())),
                len([item for item in self._planned_items if item[1].startswith(family)]),
                len([item for item in self.completed_items if item[1].startswith(family)]),
            ])
        detail = workbook.create_sheet("Issues")
        detail.append(columns)
        for issue in self.issues:
            detail.append([getattr(issue, column) for column in columns])
        workbook.save(summary_path)
        workbook.close()
        return summary_path

    def _summary_products_by_family(self, summary: ProductIdSummary) -> dict[str, tuple[str, ...]]:
        if summary.products_by_family:
            return {
                family: tuple(summary.products_by_family.get(family, ()))
                for family in CODE_SCOPE
            }
        selected_families = {code[0] for code in self.request.image_codes}
        if len(selected_families) == 1:
            family = next(iter(selected_families))
            return {key: tuple(summary.valid_ids) if key == family else () for key in CODE_SCOPE}
        return {family: () for family in CODE_SCOPE}

    def _normal_jobs(
        self, summary: ProductIdSummary,
    ) -> list[tuple[str, list[str], tuple[str, ...]]]:
        products_by_family = self._summary_products_by_family(summary)
        codes_by_family = {
            family: tuple(code for code in self.request.image_codes if code.startswith(family))
            for family in CODE_SCOPE
        }
        missing = [
            AOI_GROUP_BY_FAMILY[family]
            for family, codes in codes_by_family.items()
            if codes and not products_by_family[family]
        ]
        if missing:
            details = "、".join(f"{group.sheet_name}（{group.station_name}）" for group in missing)
            raise ValueError(f"所选图片代码缺少对应AOI产品号：{details}")

        self._products_by_family = {
            family: products_by_family[family]
            for family, codes in codes_by_family.items() if codes
        }
        self._sources_by_family = {
            family: summary.sources_by_family.get(family, "兼容输入")
            for family in self._products_by_family
        }
        self._all_product_ids = list(dict.fromkeys(
            product_id
            for family in CODE_SCOPE
            for product_id in self._products_by_family.get(family, ())
        ))
        self._planned_items = {
            (product_id, code)
            for family, product_ids in self._products_by_family.items()
            for product_id in product_ids
            for code in codes_by_family[family]
        }
        jobs: list[tuple[str, list[str], tuple[str, ...]]] = []
        for family in CODE_SCOPE:
            product_ids = self._products_by_family.get(family, ())
            codes = codes_by_family[family]
            for index in range(0, len(product_ids), self.request.batch_size):
                jobs.append((family, list(product_ids[index:index + self.request.batch_size]), codes))
        return jobs

    def _retry_jobs(self, summary: ProductIdSummary) -> list[tuple[str, list[str], tuple[str, ...]]]:
        payload = json.loads(self.request.retry_manifest.read_text(encoding="utf-8"))
        selected_codes = set(self.request.image_codes)
        failed_pairs = {
            (str(item.get("product_id", "")).strip().upper(), str(item.get("image_code", "")).strip().upper())
            for item in payload.get("failed_items", [])
            if str(item.get("product_id", "")).strip() and str(item.get("image_code", "")).strip().upper() in selected_codes
        }
        if not failed_pairs:
            raise ValueError("异常任务清单中没有可重试的产品号—图片代码项")
        self.completed_items.update(
            (str(item[0]).upper(), str(item[1]).upper())
            for item in payload.get("completed_items", []) if len(item) == 2
        )
        stored_planned = {
            (str(item[0]).upper(), str(item[1]).upper())
            for item in payload.get("planned_items", []) if len(item) == 2
        }
        self._planned_items = stored_planned or (self.completed_items | failed_pairs)
        stored_groups = payload.get("product_ids_by_family", {})
        if stored_groups:
            self._products_by_family = {
                family: tuple(str(item).upper() for item in stored_groups.get(family, []))
                for family in CODE_SCOPE if stored_groups.get(family)
            }
        else:
            current = self._summary_products_by_family(summary)
            self._products_by_family = {
                family: tuple(dict.fromkeys(
                    list(current.get(family, ()))
                    + [product_id for product_id, code in failed_pairs if code.startswith(family)]
                ))
                for family in CODE_SCOPE
                if current.get(family) or any(code.startswith(family) for _, code in failed_pairs)
            }
        self._sources_by_family = {
            **summary.sources_by_family,
            **{str(key): str(value) for key, value in payload.get("sources_by_family", {}).items()},
        }
        stored_products = [str(item).upper() for item in payload.get("product_ids", [])]
        self._all_product_ids = list(dict.fromkeys(
            stored_products or [item for products in self._products_by_family.values() for item in products]
        ))
        # 精确按失败项重试，避免把稀疏失败组合重新扩展为笛卡尔积。
        return [(code[0], [product_id], (code,)) for product_id, code in sorted(failed_pairs)]

    def run(self, product_summary: ProductIdSummary | None = None) -> ImageDownloadResult:
        self.request.validate()
        summary = product_summary or extract_product_ids(self.request.workbook_path)
        jobs = self._retry_jobs(summary) if self.request.retry_manifest else self._normal_jobs(summary)
        if not jobs:
            raise ValueError("MES工作簿中没有可下载的AOI产品号")
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.temp_dir.mkdir(parents=True, exist_ok=True)
        for scope in CODE_SCOPE.values():
            (self.output_dir / "images" / scope).mkdir(parents=True, exist_ok=True)
        self._checkpoint()
        for family, product_ids in self._products_by_family.items():
            group = AOI_GROUP_BY_FAMILY[family]
            self.log(
                f"{group.sheet_name}/{group.station_name}：有效产品 {len(product_ids)} 个，"
                f"来源 {self._sources_by_family.get(family, '未知')}"
            )
        try:
            total_batches = len(jobs)
            for batch_number, (_family, product_ids, image_codes) in enumerate(jobs, 1):
                self._check_cancelled()
                self.progress(
                    int((batch_number - 1) / max(1, total_batches) * 90),
                    f"正在处理第 {batch_number}/{total_batches} 批",
                )
                self._download_or_isolate(product_ids, image_codes)
                self._record_manifest()
            self.progress(94, "正在生成下载报告")
            selected_families = {code[0] for code in self.request.image_codes}
            if not self.request.retry_manifest:
                for family in selected_families:
                    for invalid in summary.invalid_ids_by_family.get(family, ()):
                        self.issues.append(ImageDownloadIssue(
                            invalid, "", "input", "invalid_product_id",
                            f"{AOI_GROUP_BY_FAMILY[family].sheet_name}中的产品号不是25位",
                            status="failed",
                        ))
            self._record_manifest()
            summary_path = self._write_reports()
            failed_count = len([issue for issue in self.issues if issue.status == "failed"])
            status = "partial" if failed_count else "complete"
            self.progress(100, "图片下载完成" if status == "complete" else "图片下载部分完成")
            roots = {
                scope: self.output_dir / "images" / scope
                for scope in CODE_SCOPE.values()
                if any((self.output_dir / "images" / scope).iterdir())
            }
            return ImageDownloadResult(
                status, self.output_dir, roots, len(self._all_product_ids), len(self.completed_items),
                self.issues, self.manifest_path, summary_path,
            )
        finally:
            self.backend.close()


def default_image_download_dir(output_root: Path) -> Path:
    return output_root / f"image_download_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
